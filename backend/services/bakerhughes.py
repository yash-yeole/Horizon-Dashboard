"""Baker Hughes Rig Count — free official Excel files (used with attribution).

We scrape the server-rendered rig-count pages for the *current* file links, then
download and parse the official Excel workbooks:
  - North America (weekly): totals, oil/gas, trajectory, basin, location
  - Worldwide / International (monthly): by region, worldwide total, history

No paid API, no headless browser. Source: Baker Hughes (rigcount.bakerhughes.com).
"""
from __future__ import annotations

import asyncio
import io
import re
from datetime import datetime, timezone

import httpx
import openpyxl

from config import settings
from models import RigCountResponse, RigGroup, RigItem, RigSeries, RigSeriesPoint

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
_TOTAL_LABELS = {"united states", "canada", "north america", "international", "worldwide"}
_REGIONS = {"Latin America", "Europe", "Africa", "Middle East", "Asia-Pacific"}
_MONTHS = {m: i for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july",
     "august", "september", "october", "november", "december"], start=1)}


class BakerHughesError(RuntimeError):
    pass


def _f(v) -> float | None:
    try:
        return round(float(v), 2) if v is not None else None
    except (TypeError, ValueError):
        return None


def _is_date(v) -> bool:
    if isinstance(v, datetime):
        return True
    return isinstance(v, str) and bool(re.search(r"\d{1,2}[/-]", v))


def _na_date(filename: str) -> str:
    m = re.search(r"(\d{2})-(\d{2})-(\d{4})", filename)
    return f"{m.group(3)}-{m.group(1)}-{m.group(2)}" if m else ""


def _ww_date(filename: str) -> str:
    m = re.search(r"([A-Za-z]+)[-\s]+(\d{4})", filename)
    if not m:
        return ""
    mo = _MONTHS.get(m.group(1).lower())
    return f"{m.group(2)}-{mo:02d}" if mo else ""


# ---------------- discovery + download ----------------

async def _find_file(client: httpx.AsyncClient, page: str, pattern: str) -> tuple[str, str]:
    html = (await client.get(settings.bh_base_url + page)).text
    links = list(dict.fromkeys(re.findall(r"static-files/[0-9a-f-]{36}", html)))
    rx = re.compile(pattern, re.I)
    for link in links:
        url = settings.bh_base_url + link
        async with client.stream("GET", url) as r:
            cd = r.headers.get("content-disposition", "")
            if rx.search(cd):
                fn = re.search(r'filename="?([^";]+)', cd)
                return url, (fn.group(1) if fn else "")
    raise BakerHughesError(f"current file not found on {page}")


async def _download(client: httpx.AsyncClient, url: str) -> bytes:
    r = await client.get(url)
    r.raise_for_status()
    return r.content


# ---------------- North America parsing ----------------

def _section(rows: list, header: str) -> list[RigItem]:
    """Items of the first group under a Breakdown section header (US group)."""
    items: list[RigItem] = []
    capturing = False
    for r in rows:
        label = str(r[1]).strip() if r[1] is not None else ""
        if not capturing:
            if label.lower() == header.lower() and _is_date(r[2]):
                capturing = True
            continue
        if not label:
            if items:
                break
            continue
        if label.lower() in _TOTAL_LABELS:
            break
        val = _f(r[2])
        if val is None:
            continue
        items.append(RigItem(label=label, value=val, change=_f(r[5]), year_ago=_f(r[9])))
    return items


def _na_totals(rows: list) -> list[RigItem]:
    want = ["United States", "Canada", "North America"]
    out: dict[str, RigItem] = {}
    for r in rows:
        label = str(r[1]).strip() if r[1] is not None else ""
        if label in want and label not in out:
            val = _f(r[2])
            if val is not None:
                out[label] = RigItem(label=label, value=val, change=_f(r[5]), year_ago=_f(r[9]))
    return [out[k] for k in want if k in out]


def _parse_na(data: bytes) -> tuple[list[RigItem], list[RigGroup]]:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb["NAM Breakdown"]
        rows = list(ws.iter_rows(min_row=1, max_row=100, max_col=10, values_only=True))
    finally:
        wb.close()

    summary = _na_totals(rows)
    groups = []
    for name, header in [("Oil vs Gas", "DrillFor"), ("Trajectory", "Trajectory"),
                         ("Basin", "Basin"), ("Location", "Location")]:
        items = _section(rows, header)
        if items:
            groups.append(RigGroup(name=name, items=items))
    return summary, groups


# ---------------- Worldwide parsing ----------------

def _parse_ww(data: bytes) -> tuple[list[RigItem], list[RigItem], list[RigSeries]]:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        regions: list[RigItem] = []
        intl_summary: list[RigItem] = []

        for r in wb["WW Summary"].iter_rows(min_row=1, max_row=25, max_col=9, values_only=True):
            label = str(r[1]).strip() if r[1] is not None else ""
            total, var = _f(r[5]), _f(r[7])
            if label in _REGIONS and total is not None:
                regions.append(RigItem(label=label, value=total, change=var))
            elif label == "International" and total is not None:
                intl_summary.append(RigItem(label="International", value=total, change=var))

        for r in wb["WW Breakdown"].iter_rows(min_row=1, max_row=30, max_col=8, values_only=True):
            label = str(r[1]).strip() if r[1] is not None else ""
            if label == "Worldwide":
                val = _f(r[2])
                if val is not None:
                    intl_summary.append(RigItem(label="Worldwide", value=val, change=_f(r[5])))
                break

        history = _ww_history(wb["WW Monthly"])
    finally:
        wb.close()
    return intl_summary, regions, history


def _ww_history(ws, months: int = 24) -> list[RigSeries]:
    # locate the normalized header row
    cols: dict[str, int] = {}
    header_row = 0
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=20, max_col=12, values_only=True), start=1):
        vals = [str(v).strip() if v is not None else "" for v in r]
        if "Rig Count Value" in vals and "Region" in vals and "Year" in vals and "Month" in vals:
            cols = {v: j for j, v in enumerate(vals)}
            header_row = i
            break
    if not header_row:
        return []

    ri, yi, mi, vi = cols["Region"], cols["Year"], cols["Month"], cols["Rig Count Value"]
    agg: dict[tuple[int, int], dict[str, float]] = {}
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        try:
            region = str(r[ri]).strip()
            y, m, v = int(r[yi]), int(r[mi]), float(r[vi])
        except (TypeError, ValueError, IndexError):
            continue
        bucket = agg.setdefault((y, m), {})
        bucket[region] = bucket.get(region, 0.0) + v
        bucket["Worldwide"] = bucket.get("Worldwide", 0.0) + v

    keys = sorted(agg.keys())[-months:]
    out: list[RigSeries] = []
    for name in ["Worldwide", "Middle East", "Asia-Pacific", "Latin America", "Europe", "Africa", "North America"]:
        pts = [RigSeriesPoint(date=f"{y}-{m:02d}", value=round(agg[(y, m)].get(name, 0.0), 1)) for (y, m) in keys]
        if any(p.value for p in pts):
            out.append(RigSeries(name=name, points=pts))
    return out


# ---------------- top level ----------------

async def fetch_rigcount() -> RigCountResponse:
    async with httpx.AsyncClient(timeout=120, follow_redirects=True, headers=UA) as client:
        na_url, na_name = await _find_file(client, settings.bh_na_page, r"North.America Rig.?Count Report\.xlsx")
        ww_url, ww_name = await _find_file(client, settings.bh_intl_page, r"WorldWide Rig Count Report.*\.xlsx")
        na_data, ww_data = await asyncio.gather(_download(client, na_url), _download(client, ww_url))

    na_summary, na_groups = _parse_na(na_data)
    intl_summary, intl_regions, ww_history = _parse_ww(ww_data)

    return RigCountResponse(
        na_report_date=_na_date(na_name),
        ww_report_date=_ww_date(ww_name),
        fetched_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
        stale=False,
        na_summary=na_summary,
        na_groups=na_groups,
        intl_summary=intl_summary,
        intl_regions=intl_regions,
        ww_history=ww_history,
    )
